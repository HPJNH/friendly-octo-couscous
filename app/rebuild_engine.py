from __future__ import annotations

import json
import hashlib
import re
import sqlite3
from collections import defaultdict
from copy import deepcopy
from datetime import date, datetime, timedelta
from difflib import SequenceMatcher
from pathlib import Path

from docx import Document
from flask import current_app, has_app_context

from .constants import (
    BUSINESS_TAG_CANONICAL_MAP,
    DRAFT_STANDARD_SUBSECTIONS,
    DRAFT_STRUCTURED_TABLE_COLUMNS,
    DRAFT_STRUCTURED_TABLE_COLUMNS_V2,
    LEGACY_FUZZY_LINK_BACKFILL_ENABLED,
    LEGACY_FUZZY_LINK_BACKFILL_THRESHOLD,
    LEGACY_V1_FUZZY_MATCH_ENABLED,
    LEGACY_V1_FUZZY_MATCH_THRESHOLD,
    PARSER_VERSION,
    SECTION_DEFINITIONS,
    SECTION_SUBCATEGORY_RULES,
)
from .db import get_connection
from .parsers import parse_draft_file
from .rendering import build_card, choose_section_template, template_label
from .utils import detect_date_from_filename, load_json, normalize_compare_text, now_string


DISPLAY_STATUS_ORDER = ["新增", "更新", "背景补充", "历史保留", "占位项", "无内容"]
ENTRY_SOURCE_LEVELS = {"A1": 4, "A2": 3, "B": 2, "C": 1, "占位": 0}
LEGACY_FIVE_COLUMN_HEADERS = ["标题", "时间", "来源", "核心内容", "为什么值得纳入"]
PLACEHOLDER_PATTERNS = [
    "未监测到有效新增信号",
    "未监测到额外重点新信号",
    "未发现合格内容",
    "未启用背景补充",
    "仍未发现合格内容",
    "暂无新增",
    "暂无重点",
]
WEAK_SOURCE_PATTERNS = ["搜狐", "自媒体", "公开内容推广", "行业分析", "聚合", "普通用户", "转载"]
OFFICIAL_SOURCE_PATTERNS = ["政府", "政务", "官网", "官方", "官微", "官号", "博物馆", "协会", "人民日报", "新华社"]
SOURCE_URL_PATTERN = re.compile(r"https?://\S+", re.I)
DATE_PATTERN = re.compile(r"(20\d{2})[.\-/年](\d{1,2})[.\-/月](\d{1,2})")
CHINESE_DATE_PATTERN = re.compile(r"(20\d{2})年(\d{1,2})月(\d{1,2})日")
SECTION_TYPE_BY_SUFFIX = {"1": "新增", "2": "重点", "3": "背景", "4": "当前状态"}


def rebuild_repaired_entries(mode: str = "full", write_reports: bool = True) -> dict:
    base_dir = current_app.config["BASE_DIR"]
    database_path = current_app.config["DATABASE_PATH"]
    started_at = now_string()
    run_key = f"{datetime.now().strftime('%Y%m%d-%H%M%S')}-{mode}"

    with get_connection(database_path) as connection:
        ensure_rebuild_tables(connection)
        run_id = connection.execute(
            """
            INSERT INTO rebuild_runs (
                run_key, started_at, parser_version, status
            ) VALUES (?, ?, ?, 'running')
            """,
            (run_key, started_at, PARSER_VERSION),
        ).lastrowid
        connection.commit()

    evidence_sources = scan_repair_sources(base_dir)
    linked_registry = build_linked_source_registry(evidence_sources)
    active_drafts = load_active_current_drafts(database_path)
    raw_entries: list[dict] = []
    extraction_notes: list[dict] = []
    for draft_row in active_drafts:
        payload = load_draft_payload_from_row(draft_row)
        extracted, notes = extract_entries_from_payload(draft_row, payload)
        raw_entries.extend(extracted)
        extraction_notes.extend(notes)

    normalized_entries, decisions = normalize_entry_states(raw_entries)
    deduped_entries, dedupe_decisions = dedupe_events(normalized_entries)
    decisions.extend(dedupe_decisions)
    annotate_entry_timelines(deduped_entries)
    enrich_entries_with_linked_sources(deduped_entries, linked_registry)

    source_index_path = evidence_map_path = decisions_path = migration_report_path = None
    status_check_path = None
    if write_reports:
        source_index_path = write_source_index(base_dir, evidence_sources)
        evidence_map_path = write_evidence_map(base_dir, evidence_sources, extraction_notes, decisions)
        decisions_path = write_rebuild_decisions(base_dir, deduped_entries, decisions)
        migration_report_path = write_migration_report(
            base_dir,
            active_drafts,
            evidence_sources,
            extraction_notes,
            deduped_entries,
            decisions,
            mode,
        )
        status_check_path = write_system_status_check(base_dir, active_drafts, deduped_entries)

    summary = persist_rebuilt_entries(
        database_path=database_path,
        run_id=run_id,
        entries=deduped_entries,
        source_index_path=source_index_path,
        evidence_map_path=evidence_map_path,
        decisions_path=decisions_path,
        migration_report_path=migration_report_path,
    )
    summary["mode"] = mode
    summary["source_index_path"] = str(source_index_path) if source_index_path else ""
    summary["evidence_map_path"] = str(evidence_map_path) if evidence_map_path else ""
    summary["decisions_path"] = str(decisions_path) if decisions_path else ""
    summary["migration_report_path"] = str(migration_report_path) if migration_report_path else ""
    summary["status_check_path"] = str(status_check_path) if status_check_path else ""
    summary["run_key"] = run_key
    summary["started_at"] = started_at
    summary["completed_at"] = now_string()

    with get_connection(database_path) as connection:
        connection.execute(
            """
            UPDATE rebuild_runs
            SET completed_at = ?, status = 'completed', source_index_path = ?, evidence_map_path = ?,
                decisions_path = ?, migration_report_path = ?, summary_json = ?
            WHERE id = ?
            """,
            (
                summary["completed_at"],
                str(source_index_path) if source_index_path else "",
                str(evidence_map_path) if evidence_map_path else "",
                str(decisions_path) if decisions_path else "",
                str(migration_report_path) if migration_report_path else "",
                json.dumps(summary, ensure_ascii=False),
                run_id,
            ),
        )
        connection.commit()

    return summary


def normalize_existing_entries(write_reports: bool = True) -> dict:
    return rebuild_repaired_entries(mode="normalize", write_reports=write_reports)


def dedupe_existing_entries(write_reports: bool = True) -> dict:
    return rebuild_repaired_entries(mode="dedupe", write_reports=write_reports)


def ensure_rebuild_tables(connection: sqlite3.Connection) -> None:
    connection.executescript(
        """
        CREATE TABLE IF NOT EXISTS rebuild_runs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            run_key TEXT NOT NULL,
            started_at TEXT NOT NULL,
            completed_at TEXT,
            parser_version TEXT NOT NULL,
            status TEXT NOT NULL,
            source_index_path TEXT,
            evidence_map_path TEXT,
            decisions_path TEXT,
            migration_report_path TEXT,
            summary_json TEXT
        );

        CREATE TABLE IF NOT EXISTS entries (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            run_id INTEGER,
            origin_document_id INTEGER,
            report_date TEXT NOT NULL,
            module_id INTEGER NOT NULL,
            module_key TEXT NOT NULL,
            module_name TEXT NOT NULL,
            subsection_path TEXT NOT NULL,
            subsection_title TEXT NOT NULL,
            section_type TEXT NOT NULL,
            source_level TEXT NOT NULL,
            entry_type TEXT NOT NULL,
            event_key TEXT NOT NULL,
            title TEXT NOT NULL,
            time_text TEXT,
            event_date TEXT,
            source_name TEXT,
            source_url TEXT,
            core_content TEXT NOT NULL,
            why_included TEXT,
            note_text TEXT,
            is_in_patch_window INTEGER NOT NULL DEFAULT 0,
            is_in_focus_window INTEGER NOT NULL DEFAULT 0,
            display_status TEXT NOT NULL,
            needs_review INTEGER NOT NULL DEFAULT 0,
            confidence_level TEXT NOT NULL DEFAULT '中',
            is_current_chain INTEGER NOT NULL DEFAULT 1,
            is_deleted INTEGER NOT NULL DEFAULT 0,
            dedupe_rank INTEGER NOT NULL DEFAULT 0,
            evidence_json TEXT,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        );
        """
    )


def scan_repair_sources(base_dir: Path) -> list[dict]:
    evidence_roots = [
        current_app.config["RAW_DATA_ROOT"],
        current_app.config["REVIEW_DATA_ROOT"],
        current_app.config["VERIFICATION_DATA_ROOT"],
        current_app.config["DOCS_ROOT"],
        current_app.config["EXPORT_ROOT"],
        current_app.config["REPORT_EXPORT_ROOT"],
        current_app.config["ARCHIVE_ROOT"],
    ]
    seen_paths: set[str] = set()
    items: list[dict] = []
    for root in evidence_roots:
        if not root.exists():
            continue
        for path in root.rglob("*"):
            if not path.is_file():
                continue
            if path.suffix.lower() not in {".docx", ".pdf", ".md", ".txt", ".json"}:
                continue
            if "__pycache__" in path.parts:
                continue
            path_key = str(path.resolve())
            if path_key in seen_paths:
                continue
            seen_paths.add(path_key)
            info = classify_evidence_file(base_dir, path)
            if info:
                items.append(info)
    items.sort(key=lambda item: (item.get("report_date") or "9999-99-99", item["purpose"], item["name"]))
    return items


def classify_evidence_file(base_dir: Path, path: Path) -> dict | None:
    relative_path = path.resolve().relative_to(base_dir.parent.resolve())
    file_name = path.name
    lowered = file_name.lower()
    purpose = ""
    file_type = path.suffix.lower().lstrip(".")
    if "核验" in file_name or "真实性" in file_name:
        purpose = "核验依据"
    elif "复查" in file_name:
        purpose = "复查依据"
    elif "模板" in file_name or "spec" in lowered:
        purpose = "模板规范"
    elif path.suffix.lower() == ".pdf" and ("浏览成果" in file_name or "pdf" in lowered):
        purpose = "展示结果"
    elif "简报" in file_name:
        purpose = "简报源"
    elif "底稿" in file_name or "研究底稿" in file_name:
        purpose = "底稿源"
    elif "error" in lowered or "失败" in file_name or "报错" in file_name:
        purpose = "报错依据"
    elif path.suffix.lower() == ".json" and "archive" in str(path):
        purpose = "解析归档"
    else:
        return None

    report_date = detect_date_from_filename(file_name) or find_date_in_parts(path.parts) or ""
    return {
        "name": file_name,
        "path": str(path.resolve()),
        "relative_path": str(relative_path),
        "file_type": file_type,
        "report_date": report_date,
        "version_hint": build_version_hint(file_name, report_date),
        "purpose": purpose,
        "folder": str(path.parent.resolve()),
    }


def find_date_in_parts(parts: tuple[str, ...]) -> str | None:
    for part in parts:
        if re.fullmatch(r"20\d{2}-\d{2}-\d{2}", part):
            return part
    return None


def build_version_hint(file_name: str, report_date: str | None) -> str:
    if report_date:
        return report_date
    match = re.search(r"(\d+\.\d+)", file_name)
    if match:
        return match.group(1)
    return "未标注"


def load_active_current_drafts(database_path: Path) -> list[sqlite3.Row]:
    with get_connection(database_path) as connection:
        return connection.execute(
            """
            SELECT *
            FROM documents
            WHERE doc_type = 'draft'
              AND lifecycle_status = 'active'
              AND is_current = 1
            ORDER BY report_date ASC, uploaded_at ASC
            """
        ).fetchall()


def load_draft_payload_from_row(row) -> dict:
    stored_path = Path(row["stored_path"]) if row["stored_path"] else None
    if stored_path and stored_path.exists():
        return parse_draft_file(stored_path, row["content"] or "")

    parsed_path = Path(row["parsed_path"]) if row["parsed_path"] else None
    if parsed_path and parsed_path.exists():
        payload = load_json(parsed_path.read_text(encoding="utf-8"), {})
        document_metadata = payload.get("document_metadata", {}) or {}
        return {
            "content": payload.get("content", row["content"] or ""),
            "sections": payload.get("sections", {}),
            "tail_hits": document_metadata.get("tail_hits", []),
            "metadata": document_metadata.get("draft_metadata")
            or {
                "contract_version": document_metadata.get("contract_version", "unknown"),
                "window_metadata": document_metadata.get("window_metadata", {}),
            },
        }

    return {"content": row["content"] or "", "sections": {}, "tail_hits": [], "metadata": {}}


def extract_entries_from_payload(draft_row, payload: dict) -> tuple[list[dict], list[dict]]:
    sections = payload.get("sections") or {}
    report_note_text = (sections.get("report_note", {}) or {}).get("plain_text", "")
    payload_metadata = payload.get("metadata") or {}
    patch_window, focus_window = derive_windows(draft_row["report_date"], report_note_text, payload_metadata)
    extracted: list[dict] = []
    notes: list[dict] = []

    for definition in SECTION_DEFINITIONS:
        if definition["key"] == "report_note":
            continue
        section = sections.get(definition["key"], {})
        blocks = section.get("blocks") or []
        current_path = ""
        current_title = ""
        seen_subsections: defaultdict[str, int] = defaultdict(int)
        for block in blocks:
            block_type = block.get("type")
            if block_type == "heading":
                current_path = canonical_subsection_path(block.get("path") or "")
                current_title = block.get("title") or block.get("text") or ""
                seen_subsections[current_path] += 1
                if seen_subsections[current_path] > 1:
                    notes.append(
                        {
                            "report_date": draft_row["report_date"],
                            "module_key": definition["key"],
                            "issue": "重复二级结构",
                            "message": f"{definition['title']} 出现重复结构 {current_path} {current_title}",
                            "suggestion": "保留单个固定二级结构，避免同一模块重复 2.2 / 3.2 等标题。",
                        }
                    )
                continue

            if not current_path:
                continue

            section_type = section_type_from_path(current_path)
            if not section_type:
                continue

            if block_type == "table" and section_type in {"新增", "重点", "背景"}:
                table_entries, table_notes = extract_table_entries(
                    draft_row=draft_row,
                    definition=definition,
                    subsection_path=current_path,
                    subsection_title=current_title,
                    rows=block.get("rows") or [],
                    patch_window=patch_window,
                    focus_window=focus_window,
                )
                extracted.extend(table_entries)
                notes.extend(table_notes)
                continue

            if block_type == "paragraph" and section_type == "当前状态":
                text = (block.get("text") or "").strip()
                if text:
                    extracted.append(
                        build_note_entry(
                            draft_row=draft_row,
                            definition=definition,
                            subsection_path=current_path,
                            subsection_title=current_title,
                            text=text,
                        )
                    )
                continue

            if block_type == "paragraph" and section_type in {"新增", "重点", "背景"}:
                text = (block.get("text") or "").strip()
                if not text:
                    continue
                if looks_like_placeholder_text(text):
                    extracted.append(
                        build_placeholder_entry(
                            draft_row=draft_row,
                            definition=definition,
                            subsection_path=current_path,
                            subsection_title=current_title,
                            placeholder_text=text,
                            patch_window=patch_window,
                            focus_window=focus_window,
                            salvage_mode=True,
                        )
                    )
                else:
                    notes.append(
                        {
                            "report_date": draft_row["report_date"],
                            "module_key": definition["key"],
                            "issue": "结构区自由文本",
                            "message": f"{definition['title']} 的 {current_path} 出现自由文本：{text[:48]}",
                            "suggestion": "结构化区必须改成 v1 6 列或 v2 9 列表格；本次迁移已忽略该段，以免继续放大脏数据。",
                        }
                    )
                continue

    return extracted, notes


def derive_windows(report_date_text: str, report_note_text: str, metadata: dict | None = None) -> tuple[tuple[date, date], tuple[date, date]]:
    report_day = date.fromisoformat(report_date_text)
    default_start = report_day - timedelta(days=2)
    window_metadata = (metadata or {}).get("window_metadata") or {}

    patch_window = coerce_window_range(window_metadata.get("patch_window"))
    focus_window = coerce_window_range(window_metadata.get("focus_window"))

    if not patch_window:
        discovered_dates = extract_all_dates(report_note_text)
        if len(discovered_dates) >= 2:
            patch_window = (min(discovered_dates), max(discovered_dates))
        else:
            patch_window = (default_start, report_day)

    if not focus_window:
        focus_window = (report_day - timedelta(days=2), report_day)

    return patch_window, focus_window


def coerce_window_range(raw_value) -> tuple[date, date] | None:
    if not raw_value:
        return None
    if isinstance(raw_value, dict):
        values = [raw_value.get("start"), raw_value.get("end")]
    elif isinstance(raw_value, (list, tuple)):
        values = list(raw_value)[:2]
    else:
        return None
    if len(values) < 2 or not values[0] or not values[1]:
        return None
    try:
        start = date.fromisoformat(str(values[0]))
        end = date.fromisoformat(str(values[1]))
    except ValueError:
        return None
    return (min(start, end), max(start, end))


def extract_all_dates(text: str) -> list[date]:
    values: list[date] = []
    for pattern in (DATE_PATTERN, CHINESE_DATE_PATTERN):
        for year, month, day in pattern.findall(text or ""):
            try:
                values.append(date(int(year), int(month), int(day)))
            except ValueError:
                continue
    return values


def extract_first_date(text: str) -> date | None:
    values = extract_all_dates(text or "")
    return values[0] if values else None


def extract_table_entries(
    draft_row,
    definition: dict,
    subsection_path: str,
    subsection_title: str,
    rows: list[list[str]],
    patch_window: tuple[date, date],
    focus_window: tuple[date, date],
) -> tuple[list[dict], list[dict]]:
    normalized_rows = [trim_row(row) for row in rows if any((cell or "").strip() for cell in row)]
    if not normalized_rows:
        return [], []

    headers = normalized_rows[0]
    body_rows = normalized_rows[1:]
    notes: list[dict] = []
    header_mode = identify_table_header_mode(headers)
    if header_mode is None:
        notes.append(
            {
                "report_date": draft_row["report_date"],
                "module_key": definition["key"],
                "issue": "表格列数不符",
                "message": f"{definition['title']} 的 {subsection_path} 表格列名为 {' / '.join(headers)}",
                "suggestion": "固定使用 v1 6 列或 v2 9 列模板；v2 需额外包含原始链接 / 本次新增事实 / 业务标签。",
            }
        )
        return [], notes

    extracted: list[dict] = []
    for row_index, row in enumerate(body_rows, start=1):
        mapped = map_row_to_entry_fields(headers, row, header_mode)
        if not any(mapped.values()):
            continue
        if looks_like_placeholder_row(mapped):
            extracted.append(
                build_placeholder_entry(
                    draft_row=draft_row,
                    definition=definition,
                    subsection_path=subsection_path,
                    subsection_title=subsection_title,
                    placeholder_text=mapped.get("title") or mapped.get("core_content") or "本模块无有效内容",
                    patch_window=patch_window,
                    focus_window=focus_window,
                    salvage_mode=(header_mode == "legacy_5"),
                )
            )
            continue

        extracted.append(
            build_real_entry(
                draft_row=draft_row,
                definition=definition,
                subsection_path=subsection_path,
                subsection_title=subsection_title,
                mapped=mapped,
                patch_window=patch_window,
                focus_window=focus_window,
                legacy_mode=(header_mode == "legacy_5"),
                contract_version="v2" if header_mode == "contract_9" else ("v1" if header_mode == "contract_6" else "legacy_5"),
                row_index=row_index,
            )
        )
        if header_mode == "legacy_5":
            notes.append(
                {
                    "report_date": draft_row["report_date"],
                    "module_key": definition["key"],
                    "issue": "历史 5 列表格已迁移",
                    "message": f"{definition['title']} 的 {subsection_path} 第 {row_index} 行来自旧版 5 列结构，已自动补推来源层级。",
                    "suggestion": "后续新底稿请统一改为 6 列正式模板，避免再次依赖迁移兜底。",
                }
            )
    return extracted, notes


def identify_table_header_mode(headers: list[str]) -> str | None:
    compact_headers = [item.strip() for item in headers]
    if compact_headers == DRAFT_STRUCTURED_TABLE_COLUMNS:
        return "contract_6"
    if compact_headers == DRAFT_STRUCTURED_TABLE_COLUMNS_V2:
        return "contract_9"
    if compact_headers == LEGACY_FIVE_COLUMN_HEADERS:
        return "legacy_5"
    return None


def map_row_to_entry_fields(headers: list[str], row: list[str], header_mode: str) -> dict:
    values = [row[index].strip() if index < len(row) else "" for index in range(len(headers))]
    if header_mode == "contract_6":
        return {
            "title": values[0],
            "time_text": values[1],
            "source_level": values[2],
            "source_name": values[3],
            "core_content": values[4],
            "why_included": values[5],
            "source_url": "",
            "delta_text": "",
            "business_tags": [],
        }
    if header_mode == "contract_9":
        return {
            "title": values[0],
            "time_text": values[1],
            "source_level": values[2],
            "source_name": values[3],
            "source_url": values[4],
            "core_content": values[5],
            "delta_text": values[6],
            "business_tags": normalize_business_tags(values[7]),
            "why_included": values[8],
        }
    return {
        "title": values[0],
        "time_text": values[1],
        "source_level": "",
        "source_name": values[2],
        "core_content": values[3],
        "why_included": values[4],
        "source_url": "",
        "delta_text": "",
        "business_tags": [],
    }


def build_real_entry(
    draft_row,
    definition: dict,
    subsection_path: str,
    subsection_title: str,
    mapped: dict,
    patch_window: tuple[date, date],
    focus_window: tuple[date, date],
    legacy_mode: bool,
    contract_version: str,
    row_index: int,
) -> dict:
    time_text = mapped.get("time_text", "").strip()
    source_name = mapped.get("source_name", "").strip()
    source_level = normalize_source_level(mapped.get("source_level", "").strip(), source_name)
    source_url = (mapped.get("source_url") or "").strip() or extract_source_url(source_name)
    delta_text = (mapped.get("delta_text") or "").strip()
    business_tags = normalize_business_tags(mapped.get("business_tags"))
    event_anchor = normalize_compare_text(mapped.get("event_anchor", ""))
    canonical_title = normalize_compare_text(mapped.get("canonical_title", ""))
    event_date = extract_first_date(time_text)
    event_key = provisional_event_key(
        definition["key"],
        mapped.get("title", ""),
        time_text,
        mapped.get("core_content", ""),
        source_name,
        canonical_title=canonical_title,
    )
    is_in_patch_window = is_date_in_window(event_date, patch_window)
    is_in_focus_window = is_date_in_window(event_date, focus_window)
    evidence = {
        "legacy_mode": legacy_mode,
        "source_file": draft_row["original_name"],
        "source_path": draft_row["stored_path"],
        "row_index": row_index,
        "contract_version": contract_version,
        "patch_window": [patch_window[0].isoformat(), patch_window[1].isoformat()],
        "focus_window": [focus_window[0].isoformat(), focus_window[1].isoformat()],
        "delta_text": delta_text,
        "business_tags": business_tags,
        "normalization_notes": [],
    }
    if event_anchor:
        evidence["event_anchor"] = event_anchor
    if canonical_title:
        evidence["canonical_title"] = canonical_title
    return {
        "origin_document_id": draft_row["id"],
        "report_date": draft_row["report_date"],
        "module_id": definition["number"],
        "module_key": definition["key"],
        "module_name": definition["title"],
        "subsection_path": subsection_path,
        "subsection_title": subsection_title or expected_subsection_title(subsection_path),
        "section_type": section_type_from_path(subsection_path),
        "source_level": source_level,
        "entry_type": "real",
        "event_key": event_key,
        "title": mapped.get("title", "").strip() or f"{definition['title']} 条目 {row_index}",
        "time_text": time_text,
        "event_date": event_date.isoformat() if event_date else "",
        "source_name": source_name,
        "source_title": mapped.get("title", "").strip(),
        "source_url": source_url,
        "supporting_sources_json": [],
        "core_content": mapped.get("core_content", "").strip(),
        "why_included": mapped.get("why_included", "").strip(),
        "note_text": "",
        "first_seen_date": draft_row["report_date"],
        "last_seen_date": draft_row["report_date"],
        "is_in_patch_window": 1 if is_in_patch_window else 0,
        "is_in_focus_window": 1 if is_in_focus_window else 0,
        "display_status": "历史保留",
        "needs_review": 0,
        "confidence_level": infer_confidence_level(source_level, False),
        "is_current_chain": 1,
        "is_deleted": 0,
        "dedupe_rank": 0,
        "evidence_json": evidence,
        "created_at": now_string(),
        "updated_at": now_string(),
    }


def build_placeholder_entry(
    draft_row,
    definition: dict,
    subsection_path: str,
    subsection_title: str,
    placeholder_text: str,
    patch_window: tuple[date, date],
    focus_window: tuple[date, date],
    salvage_mode: bool,
) -> dict:
    evidence = {
        "salvage_mode": salvage_mode,
        "source_file": draft_row["original_name"],
        "source_path": draft_row["stored_path"],
        "patch_window": [patch_window[0].isoformat(), patch_window[1].isoformat()],
        "focus_window": [focus_window[0].isoformat(), focus_window[1].isoformat()],
        "normalization_notes": ["合法占位项"],
    }
    return {
        "origin_document_id": draft_row["id"],
        "report_date": draft_row["report_date"],
        "module_id": definition["number"],
        "module_key": definition["key"],
        "module_name": definition["title"],
        "subsection_path": subsection_path,
        "subsection_title": subsection_title or expected_subsection_title(subsection_path),
        "section_type": section_type_from_path(subsection_path),
        "source_level": "占位",
        "entry_type": "placeholder",
        "event_key": f"{definition['key']}:{subsection_path}:placeholder:{draft_row['report_date']}",
        "title": placeholder_text.strip(),
        "time_text": "",
        "event_date": "",
        "source_name": "",
        "source_title": "",
        "source_url": "",
        "supporting_sources_json": [],
        "core_content": placeholder_text.strip(),
        "why_included": "",
        "note_text": "",
        "first_seen_date": draft_row["report_date"],
        "last_seen_date": draft_row["report_date"],
        "is_in_patch_window": 0,
        "is_in_focus_window": 0,
        "display_status": "占位项",
        "needs_review": 0,
        "confidence_level": "低",
        "is_current_chain": 1,
        "is_deleted": 0,
        "dedupe_rank": 0,
        "evidence_json": evidence,
        "created_at": now_string(),
        "updated_at": now_string(),
    }


def build_note_entry(draft_row, definition: dict, subsection_path: str, subsection_title: str, text: str) -> dict:
    return {
        "origin_document_id": draft_row["id"],
        "report_date": draft_row["report_date"],
        "module_id": definition["number"],
        "module_key": definition["key"],
        "module_name": definition["title"],
        "subsection_path": subsection_path,
        "subsection_title": subsection_title or expected_subsection_title(subsection_path),
        "section_type": "当前状态",
        "source_level": "A2",
        "entry_type": "real",
        "event_key": f"{definition['key']}:{subsection_path}:{normalize_compare_text(text)[:24]}:{draft_row['report_date']}",
        "title": subsection_title or "当前状态说明",
        "time_text": "",
        "event_date": "",
        "source_name": draft_row["original_name"],
        "source_title": subsection_title or "当前状态说明",
        "source_url": "",
        "supporting_sources_json": [],
        "core_content": text.strip(),
        "why_included": "",
        "note_text": text.strip(),
        "first_seen_date": draft_row["report_date"],
        "last_seen_date": draft_row["report_date"],
        "is_in_patch_window": 0,
        "is_in_focus_window": 0,
        "display_status": "历史保留",
        "needs_review": 0,
        "confidence_level": "中",
        "is_current_chain": 1,
        "is_deleted": 0,
        "dedupe_rank": 0,
        "evidence_json": {
            "source_file": draft_row["original_name"],
            "source_path": draft_row["stored_path"],
            "normalization_notes": ["当前状态说明"],
        },
        "created_at": now_string(),
        "updated_at": now_string(),
    }


def normalize_entry_states(records: list[dict]) -> tuple[list[dict], list[dict]]:
    decisions: list[dict] = []
    normalized: list[dict] = []
    previous_valid_by_module: defaultdict[str, list[dict]] = defaultdict(list)

    for record in sorted(records, key=lambda item: (item["report_date"], item["module_id"], item["subsection_path"], item["title"])):
        item = deepcopy(record)
        evidence = deepcopy(item.get("evidence_json") or {})
        notes = evidence.setdefault("normalization_notes", [])
        section_type = item["section_type"]

        if item["entry_type"] == "placeholder":
            item["display_status"] = "占位项"
            item["confidence_level"] = "低"
            normalized.append(item)
            continue

        weak_source = has_weak_source(item["source_name"])
        if weak_source:
            item["needs_review"] = 1
            notes.append("弱来源")
        if is_missing_minimum_fields(item):
            item["display_status"] = "删除"
            item["is_deleted"] = 1
            item["needs_review"] = 1
            item["confidence_level"] = "低"
            notes.append("缺少必要字段")
            decisions.append(build_decision(item, "删除", "条目缺少标题/时间/来源/核心内容中的关键字段"))
            normalized.append(item)
            continue

        previous_match = find_previous_match(previous_valid_by_module[item["module_key"]], item)
        if section_type == "背景":
            item["display_status"] = "背景补充"
        elif section_type == "当前状态":
            item["display_status"] = "历史保留"
        elif section_type == "新增":
            if not item["is_in_patch_window"]:
                item["display_status"] = "历史保留"
                notes.append("不在补采窗口")
                decisions.append(build_decision(item, "改状态", "原位于新增区，但事件时间不在补采窗口内，改为历史保留"))
            elif previous_match:
                if has_meaningful_delta(previous_match, item):
                    item["display_status"] = "更新"
                    notes.append("与上一有效版本为同一事件但存在新增事实")
                else:
                    item["display_status"] = "历史保留"
                    notes.append("同事件重复表述，不再作为新增")
                    decisions.append(build_decision(item, "改状态", "同一事件换说法重复出现，改为历史保留"))
            else:
                item["display_status"] = "新增"
        elif section_type == "重点":
            if not item["is_in_focus_window"]:
                item["display_status"] = "历史保留"
                notes.append("不在近72小时重点窗口")
                decisions.append(build_decision(item, "改状态", "原位于重点区，但事件时间不在近72小时窗口内，改为历史保留"))
            elif previous_match:
                item["display_status"] = "更新" if has_meaningful_delta(previous_match, item) else "历史保留"
                if item["display_status"] == "历史保留":
                    decisions.append(build_decision(item, "改状态", "重点区重复事件未产生新增事实，改为历史保留"))
            else:
                item["display_status"] = "新增"
        else:
            item["display_status"] = "历史保留"

        if weak_source and item["display_status"] in {"新增", "更新"}:
            item["needs_review"] = 1
            notes.append("弱来源承载新增/更新结论")
            decisions.append(build_decision(item, "待复核", "弱来源条目继续保留，但已标记 needs_review"))

        if has_overclaimed_value(item):
            item["needs_review"] = 1
            notes.append("观察价值可能超出事实支撑")
            decisions.append(build_decision(item, "待复核", "观察价值存在过度推断风险，已标记 needs_review"))

        item["confidence_level"] = infer_confidence_level(item["source_level"], bool(item["needs_review"]))
        item["evidence_json"] = evidence
        normalized.append(item)
        if item["display_status"] != "删除":
            previous_valid_by_module[item["module_key"]].append(item)

    return normalized, decisions


def dedupe_events(records: list[dict]) -> tuple[list[dict], list[dict]]:
    decisions: list[dict] = []
    items = [deepcopy(record) for record in records]
    real_items = [item for item in items if item["entry_type"] == "real" and not item["is_deleted"]]

    canonical_groups: list[dict] = []
    for item in sorted(real_items, key=lambda record: (record["report_date"], record["module_id"], record["title"])):
        candidate = find_canonical_group(canonical_groups, item)
        if candidate:
            item["event_key"] = candidate["event_key"]
        else:
            canonical_groups.append(item)

    grouped: defaultdict[str, list[dict]] = defaultdict(list)
    for item in items:
        grouped[item["event_key"]].append(item)

    result: list[dict] = []
    for _event_key, group in grouped.items():
        placeholders = [item for item in group if item["entry_type"] == "placeholder"]
        real_group = [item for item in group if item["entry_type"] == "real"]
        if not real_group:
            result.extend(placeholders)
            continue

        real_group.sort(key=entry_quality_score, reverse=True)
        canonical = real_group[0]
        canonical["dedupe_rank"] = 1
        canonical["evidence_json"] = deepcopy(canonical.get("evidence_json") or {})
        canonical["evidence_json"].setdefault("merged_from", [])

        for duplicate in real_group[1:]:
            if same_effective_fact(canonical, duplicate):
                duplicate["display_status"] = "删除"
                duplicate["is_deleted"] = 1
                decisions.append(build_decision(duplicate, "删除", "同一事件重复收录且未提供新增事实"))
                canonical["evidence_json"]["merged_from"].append(
                    {
                        "report_date": duplicate["report_date"],
                        "document_id": duplicate["origin_document_id"],
                        "title": duplicate["title"],
                    }
                )
            else:
                merge_duplicate_into_canonical(canonical, duplicate)
                duplicate["display_status"] = "删除"
                duplicate["is_deleted"] = 1
                decisions.append(build_decision(duplicate, "降权", "同事件次级来源含补充事实，已并入主条并从平铺展示中移除"))
                canonical["evidence_json"]["merged_from"].append(
                    {
                        "report_date": duplicate["report_date"],
                        "document_id": duplicate["origin_document_id"],
                        "title": duplicate["title"],
                        "merged": True,
                    }
                )

        result.append(canonical)
        result.extend(placeholders)
        result.extend(item for item in real_group[1:] if item["is_deleted"])

    result.sort(key=lambda item: (item["report_date"], item["module_id"], item["subsection_path"], item["title"]))
    return result, decisions


def build_repaired_section_views(connection: sqlite3.Connection, report_date: str, latest_draft=None) -> dict[str, dict]:
    rows = connection.execute(
        """
        SELECT e.*, d.original_name AS source_file_name, d.stored_path AS source_file_path
        FROM entries e
        LEFT JOIN documents d ON d.id = e.origin_document_id
        WHERE e.report_date <= ?
          AND e.is_current_chain = 1
        ORDER BY e.report_date ASC, e.module_id ASC, e.subsection_path ASC, e.id ASC
        """,
        (report_date,),
    ).fetchall()
    if not rows:
        return {}

    by_module: defaultdict[str, list[dict]] = defaultdict(list)
    for row in rows:
        row_map = dict(row)
        row_map["supporting_sources_json"] = load_json(row_map.get("supporting_sources_json"), [])
        row_map["evidence_json"] = load_json(row_map.get("evidence_json"), {})
        by_module[row["module_key"]].append(row_map)

    return {
        definition["key"]: build_repaired_section_rollup(definition, by_module.get(definition["key"], []), report_date, latest_draft)
        for definition in SECTION_DEFINITIONS
        if definition["key"] != "report_note"
    }


def build_repaired_section_rollup(definition: dict, rows: list[dict], report_date: str, latest_draft) -> dict:
    current_rows = [row for row in rows if row["report_date"] == report_date and not row["is_deleted"]]
    displayed_current_keys: set[str] = set()
    groups: list[dict] = []

    new_group_rows = filter_current_group_rows(current_rows, "新增")
    focus_group_rows = filter_current_group_rows(current_rows, "重点")
    background_group_rows = latest_effective_rows(rows, "背景")
    note_group_rows = latest_note_rows(rows)

    groups.extend(build_display_entry_groups(definition["key"], f"{definition['number']}.1 补采窗口内新增信号", 2, new_group_rows, "新增"))
    groups.extend(build_display_entry_groups(definition["key"], f"{definition['number']}.2 近72小时重点新信号", 2, focus_group_rows, "重点"))
    groups.extend(build_display_entry_groups(definition["key"], f"{definition['number']}.3 背景补充", 2, background_group_rows, "背景"))

    for row in new_group_rows + focus_group_rows + background_group_rows:
        if row["entry_type"] == "real" and not row["is_deleted"]:
            displayed_current_keys.add(row["event_key"])

    historical_rows = canonical_historical_rows(rows, displayed_current_keys)
    if historical_rows:
        groups.extend(build_display_entry_groups(definition["key"], "历史保留", 2, historical_rows, "历史"))

    if note_group_rows:
        groups.append(build_note_group(f"{definition['number']}.4 当前状态说明", note_group_rows))

    groups = [group for group in groups if group.get("blocks")]
    render_payload = {
        "section_key": definition["key"],
        "section_title": definition["title"],
        "template_name": choose_section_template(definition["key"]),
        "template_label": template_label(choose_section_template(definition["key"])),
        "groups": groups,
        "outline": [group["title"] for group in groups if group.get("title")],
        "card_count": sum(1 for group in groups for block in group.get("blocks", []) if block.get("type") == "card"),
        "table_count": 0,
    }
    render_payload["status_counts"] = count_group_statuses(groups)

    status = section_status_from_groups(groups, render_payload["status_counts"])
    source_dates = sorted({row["report_date"] for row in rows if not row["is_deleted"]})
    source_date = source_dates[-1] if source_dates else ""
    raw_content = render_text_from_groups(groups)
    note = build_rollup_note_from_entries(report_date, render_payload["status_counts"], source_dates)
    return {
        "section_key": definition["key"],
        "section_title": definition["title"],
        "report_date": report_date,
        "status": status,
        "note": note,
        "similarity": 1.0,
        "raw_content": raw_content,
        "display_content": raw_content,
        "metadata_json": json.dumps(
            {
                "display_render": render_payload,
                "raw_render": render_payload,
                "rollup": {"mode": "repaired_entries", "report_date": report_date},
            },
            ensure_ascii=False,
        ),
        "source_document_id": rows[-1]["origin_document_id"] if rows else (latest_draft["id"] if latest_draft else None),
        "source_date": source_date,
        "current_file_name": latest_draft["original_name"] if latest_draft else "",
        "current_file_path": latest_draft["stored_path"] if latest_draft else "",
        "document_id": rows[-1]["origin_document_id"] if rows else (latest_draft["id"] if latest_draft else None),
        "current_document_id": latest_draft["id"] if latest_draft else None,
    }


def build_repaired_section_history(connection: sqlite3.Connection, report_date: str, section_key: str, limit: int = 8) -> list[dict]:
    rows = connection.execute(
        """
        SELECT *
        FROM entries
        WHERE module_key = ?
          AND report_date <= ?
          AND is_current_chain = 1
        ORDER BY report_date DESC, subsection_path ASC, id ASC
        """,
        (section_key, report_date),
    ).fetchall()
    grouped: defaultdict[str, list[dict]] = defaultdict(list)
    for row in rows:
        grouped[row["report_date"]].append(dict(row))

    history_rows: list[dict] = []
    for day in sorted(grouped.keys(), reverse=True)[:limit]:
        day_rows = [item for item in grouped[day] if not item["is_deleted"]]
        counts = defaultdict(int)
        titles: list[str] = []
        for item in day_rows:
            counts[item["display_status"]] += 1
            if item["entry_type"] == "real" and item["title"] and item["title"] not in titles:
                titles.append(item["title"])
        history_rows.append(
            {
                "report_date": day,
                "status": section_status_from_counts(counts),
                "note": build_history_note(counts),
                "excerpt": "；".join(titles[:2]) if titles else "本日无有效结构化条目",
            }
        )
    return history_rows


def _entry_persistence_key(item: dict) -> tuple[str, str, str, str, str]:
    return (
        str(item.get("report_date") or ""),
        str(item.get("module_key") or ""),
        str(item.get("subsection_path") or ""),
        str(item.get("entry_type") or ""),
        str(item.get("event_key") or ""),
    )


def _group_existing_entries_by_key(rows: list[sqlite3.Row]) -> dict[tuple[str, str, str, str, str], list[dict]]:
    grouped: dict[tuple[str, str, str, str, str], list[dict]] = defaultdict(list)
    for row in rows:
        record = dict(row)
        grouped[_entry_persistence_key(record)].append(record)
    for bucket in grouped.values():
        bucket.sort(key=lambda item: item["id"])
    return grouped


def _upsert_rebuilt_entry(connection: sqlite3.Connection, run_id: int, item: dict, existing_row: dict | None) -> int:
    payload = (
        run_id,
        item["origin_document_id"],
        item["report_date"],
        item["module_id"],
        item["module_key"],
        item["module_name"],
        item["subsection_path"],
        item["subsection_title"],
        item["section_type"],
        item["source_level"],
        item["entry_type"],
        item["event_key"],
        item["title"],
        item.get("time_text", ""),
        item.get("event_date", ""),
        item.get("source_name", ""),
        item.get("source_title", ""),
        item.get("source_url", ""),
        json.dumps(item.get("supporting_sources_json", []), ensure_ascii=False),
        item.get("core_content", ""),
        item.get("why_included", ""),
        item.get("note_text", ""),
        item.get("first_seen_date", item["report_date"]),
        item.get("last_seen_date", item["report_date"]),
        item.get("is_in_patch_window", 0),
        item.get("is_in_focus_window", 0),
        item.get("display_status", "历史保留"),
        item.get("needs_review", 0),
        item.get("confidence_level", "中"),
        item.get("is_current_chain", 1),
        item.get("is_deleted", 0),
        item.get("dedupe_rank", 0),
        json.dumps(item.get("evidence_json", {}), ensure_ascii=False),
    )
    if existing_row:
        connection.execute(
            """
            UPDATE entries
            SET run_id = ?,
                origin_document_id = ?,
                report_date = ?,
                module_id = ?,
                module_key = ?,
                module_name = ?,
                subsection_path = ?,
                subsection_title = ?,
                section_type = ?,
                source_level = ?,
                entry_type = ?,
                event_key = ?,
                title = ?,
                time_text = ?,
                event_date = ?,
                source_name = ?,
                source_title = ?,
                source_url = ?,
                supporting_sources_json = ?,
                core_content = ?,
                why_included = ?,
                note_text = ?,
                first_seen_date = ?,
                last_seen_date = ?,
                is_in_patch_window = ?,
                is_in_focus_window = ?,
                display_status = ?,
                needs_review = ?,
                confidence_level = ?,
                is_current_chain = ?,
                is_deleted = ?,
                dedupe_rank = ?,
                evidence_json = ?,
                updated_at = ?
            WHERE id = ?
            """,
            (
                *payload,
                item.get("updated_at", now_string()),
                existing_row["id"],
            ),
        )
        return existing_row["id"]

    created_at = item.get("created_at", now_string())
    updated_at = item.get("updated_at", created_at)
    return connection.execute(
        """
        INSERT INTO entries (
            run_id, origin_document_id, report_date, module_id, module_key, module_name,
            subsection_path, subsection_title, section_type, source_level, entry_type, event_key,
            title, time_text, event_date, source_name, source_title, source_url, supporting_sources_json,
            core_content, why_included, note_text, first_seen_date, last_seen_date,
            is_in_patch_window, is_in_focus_window, display_status, needs_review,
            confidence_level, is_current_chain, is_deleted, dedupe_rank, evidence_json,
            created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (*payload, created_at, updated_at),
    ).lastrowid


def _soft_delete_unmatched_entries(connection: sqlite3.Connection, rows: list[dict]) -> None:
    if not rows:
        return
    now = now_string()
    for row in rows:
        connection.execute(
            """
            UPDATE entries
            SET is_deleted = 1,
                is_current_chain = 0,
                updated_at = ?
            WHERE id = ?
            """,
            (now, row["id"]),
        )


def _relink_entry_marks(connection: sqlite3.Connection) -> None:
    current_rows = connection.execute(
        """
        SELECT id, event_key, module_key, report_date, title
        FROM entries
        WHERE is_deleted = 0
        ORDER BY report_date DESC, id DESC
        """
    ).fetchall()
    latest_by_event_key: dict[str, sqlite3.Row] = {}
    for row in current_rows:
        event_key = str(row["event_key"] or "").strip()
        if event_key and event_key not in latest_by_event_key:
            latest_by_event_key[event_key] = row

    mark_rows = connection.execute(
        """
        SELECT id, entry_id, entry_event_key, marker_identity_id
        FROM entry_marks
        WHERE is_active = 1
        ORDER BY updated_at DESC, id DESC
        """
    ).fetchall()
    if not mark_rows:
        return

    now = now_string()
    seen_actor_event_keys: set[tuple[int, str]] = set()
    for row in mark_rows:
        event_key = str(row["entry_event_key"] or "").strip()
        dedupe_key = (int(row["marker_identity_id"] or 0), event_key)
        if event_key and dedupe_key in seen_actor_event_keys:
            connection.execute(
                """
                UPDATE entry_marks
                SET is_active = 0,
                    updated_at = ?
                WHERE id = ?
                """,
                (now, row["id"]),
            )
            continue
        if event_key:
            seen_actor_event_keys.add(dedupe_key)
        target = latest_by_event_key.get(event_key)
        if not target or row["entry_id"] == target["id"]:
            continue
        connection.execute(
            """
            UPDATE entry_marks
            SET entry_id = ?,
                entry_module_key = ?,
                entry_report_date = ?,
                entry_title = ?,
                updated_at = ?
            WHERE id = ?
            """,
            (
                target["id"],
                target["module_key"],
                target["report_date"],
                target["title"],
                now,
                row["id"],
            ),
        )


def persist_rebuilt_entries(
    database_path: Path,
    run_id: int,
    entries: list[dict],
    source_index_path: Path | None,
    evidence_map_path: Path | None,
    decisions_path: Path | None,
    migration_report_path: Path | None,
) -> dict:
    with get_connection(database_path) as connection:
        ensure_rebuild_tables(connection)
        existing_rows = connection.execute("SELECT * FROM entries ORDER BY id ASC").fetchall()
        existing_by_key = _group_existing_entries_by_key(existing_rows)
        for item in entries:
            existing_row = None
            bucket = existing_by_key.get(_entry_persistence_key(item))
            if bucket:
                existing_row = bucket.pop(0)
            _upsert_rebuilt_entry(connection, run_id, item, existing_row)

        for bucket in existing_by_key.values():
            _soft_delete_unmatched_entries(connection, bucket)

        _relink_entry_marks(connection)
        summary = build_summary_from_entries(entries)
        connection.execute(
            """
            UPDATE rebuild_runs
            SET source_index_path = ?, evidence_map_path = ?, decisions_path = ?, migration_report_path = ?, summary_json = ?
            WHERE id = ?
            """,
            (
                str(source_index_path) if source_index_path else "",
                str(evidence_map_path) if evidence_map_path else "",
                str(decisions_path) if decisions_path else "",
                str(migration_report_path) if migration_report_path else "",
                json.dumps(summary, ensure_ascii=False),
                run_id,
            ),
        )
        connection.commit()
        return summary


def build_summary_from_entries(entries: list[dict]) -> dict:
    kept_entries = [item for item in entries if not item["is_deleted"] and item["entry_type"] == "real"]
    placeholder_entries = [item for item in entries if item["entry_type"] == "placeholder" and not item["is_deleted"]]
    review_entries = [item for item in kept_entries if item.get("needs_review")]
    display_counts = defaultdict(int)
    for item in kept_entries + placeholder_entries:
        display_counts[item["display_status"]] += 1
    return {
        "entry_count": len(entries),
        "kept_real_entries": len(kept_entries),
        "placeholder_entries": len(placeholder_entries),
        "deleted_entries": sum(1 for item in entries if item["is_deleted"]),
        "needs_review_entries": len(review_entries),
        "display_counts": dict(display_counts),
    }


def write_source_index(base_dir: Path, evidence_sources: list[dict]) -> Path:
    output_path = current_app.config["REPORT_EXPORT_ROOT"] / "SOURCE_INDEX.md"
    lines = [
        "# SOURCE_INDEX",
        "",
        "本文件记录本次修复与迁移重建所扫描到的主要证据源文件，用于建立可追溯链条。",
        "",
        "| 文件名 | 文件类型 | 日期/版本 | 用途 | 原始路径 |",
        "| --- | --- | --- | --- | --- |",
    ]
    for item in evidence_sources:
        lines.append(
            f"| {item['name']} | {item['file_type']} | {item['version_hint']} | {item['purpose']} | {item['path']} |"
        )
    for path in sorted(current_app.config["LINKED_DATA_ROOT"].glob("*.docx")):
        if any(item["path"] == str(path.resolve()) for item in evidence_sources):
            continue
        lines.append(f"| {path.name} | docx | 链接资料 | 链接补全依据 | {path.resolve()} |")
    output_path.write_text("\n".join(lines), encoding="utf-8")
    return output_path


def write_evidence_map(base_dir: Path, evidence_sources: list[dict], extraction_notes: list[dict], decisions: list[dict]) -> Path:
    output_path = current_app.config["REPORT_EXPORT_ROOT"] / "EVIDENCE_MAP.md"
    verification_files = [item for item in evidence_sources if item["purpose"] == "核验依据"]
    display_files = [item for item in evidence_sources if item["purpose"] == "展示结果"]
    draft_sources = [item for item in evidence_sources if item["purpose"] == "底稿源"]
    lines = [
        "# EVIDENCE_MAP",
        "",
        "本文件记录“问题 → 证据文件 → 处理动作”的映射关系，避免只看数据库现状而忽略原始证据。",
        "",
        "## 核心问题映射",
        "",
        "| 问题 | 主要证据文件 | 处理动作 |",
        "| --- | --- | --- |",
        f"| 新增误判与窗口错位 | {join_evidence_names(verification_files[:1] + draft_sources[:2])} | 迁移时重算补采窗口 / 近72小时窗口，窗口外条目改为历史保留或删除。 |",
        f"| 结构区自由文本与占位写法不规范 | {join_evidence_names(draft_sources[:4])} | 新上传改为严格 6 列模板校验；历史自由文本仅在识别为合法占位时被迁移为占位项，其余忽略。 |",
        f"| 同一事件重复收录 | {join_evidence_names(verification_files[:1] + display_files[:1])} | 建立 event_key 与去重合并逻辑，主条保留高质量来源，次级重复条目改为删除或并入证据链。 |",
        f"| 弱来源支撑关键结论 | {join_evidence_names(verification_files[:1])} | 对弱来源条目执行降权、删除或 needs_review 标记，不再默认以新增/重点身份输出。 |",
        "",
        "## 解析与迁移备注",
        "",
    ]
    for note in extraction_notes[:20]:
        lines.append(
            f"- {note['report_date']} / {note['module_key']} / {note['issue']}：{note['message']}。处理建议：{note['suggestion']}"
        )
    if decisions:
        lines.extend(["", "## 代表性处理动作", ""])
        for decision in decisions[:20]:
            lines.append(
                f"- {decision['report_date']} / {decision['module_name']} / {decision['title']}：{decision['action']}，依据：{decision['reason']}"
            )
    output_path.write_text("\n".join(lines), encoding="utf-8")
    return output_path


def write_rebuild_decisions(base_dir: Path, entries: list[dict], decisions: list[dict]) -> Path:
    output_path = current_app.config["REPORT_EXPORT_ROOT"] / "REBUILD_DECISIONS.md"
    kept = [item for item in entries if not item["is_deleted"] and item["entry_type"] == "real"]
    downgraded = [item for item in entries if item.get("needs_review") and not item["is_deleted"]]
    deleted = [item for item in entries if item["is_deleted"]]
    placeholders = [item for item in entries if item["entry_type"] == "placeholder" and not item["is_deleted"]]
    lines = [
        "# REBUILD_DECISIONS",
        "",
        f"- 保留的真实条目：{len(kept)}",
        f"- 低可信待复核：{len(downgraded)}",
        f"- 已删除/移出展示：{len(deleted)}",
        f"- 占位项：{len(placeholders)}",
        "",
        "## 状态变更与去重决策",
        "",
    ]
    for decision in decisions:
        lines.append(
            f"- {decision['report_date']} / {decision['module_name']} / {decision['title']}：{decision['action']}。依据：{decision['reason']}"
        )
    output_path.write_text("\n".join(lines), encoding="utf-8")
    return output_path


def write_migration_report(
    base_dir: Path,
    active_drafts: list[sqlite3.Row],
    evidence_sources: list[dict],
    extraction_notes: list[dict],
    entries: list[dict],
    decisions: list[dict],
    mode: str,
) -> Path:
    output_path = current_app.config["REPORT_EXPORT_ROOT"] / "MIGRATION_REPORT.md"
    summary = build_summary_from_entries(entries)
    lines = [
        "# MIGRATION_REPORT",
        "",
        f"- 执行模式：{mode}",
        f"- 执行时间：{now_string()}",
        f"- 当前有效底稿版本数：{len(active_drafts)}",
        f"- 纳入修复依据的文件数：{len(evidence_sources)}",
        f"- 抽取出的条目总数：{summary['entry_count']}",
        f"- 保留的真实条目：{summary['kept_real_entries']}",
        f"- 占位项：{summary['placeholder_entries']}",
        f"- 已删除/移出展示条目：{summary['deleted_entries']}",
        f"- needs_review 条目：{summary['needs_review_entries']}",
        "",
        "## 展示状态统计",
        "",
    ]
    for label, count in summary["display_counts"].items():
        lines.append(f"- {label}：{count}")
    lines.extend(["", "## 历史文档链", ""])
    for row in active_drafts:
        lines.append(f"- {row['report_date']} / {row['original_name']} / 当前生效")
    if extraction_notes:
        lines.extend(["", "## 迁移过程中的结构提醒", ""])
        for note in extraction_notes[:20]:
            lines.append(f"- {note['report_date']} / {note['module_key']}：{note['issue']}，{note['message']}")
    if decisions:
        lines.extend(["", "## 关键处置摘要", ""])
        for decision in decisions[:20]:
            lines.append(f"- {decision['report_date']} / {decision['title']}：{decision['action']}，{decision['reason']}")
    output_path.write_text("\n".join(lines), encoding="utf-8")
    return output_path


def write_system_status_check(base_dir: Path, active_drafts: list[sqlite3.Row], entries: list[dict]) -> Path:
    output_path = current_app.config["REPORT_EXPORT_ROOT"] / "SYSTEM_STATUS_CHECK.md"
    by_module: defaultdict[str, defaultdict[str, int]] = defaultdict(lambda: defaultdict(int))
    for item in entries:
        if item["is_deleted"]:
            continue
        by_module[item["module_name"]][item["display_status"]] += 1
    lines = [
        "# SYSTEM_STATUS_CHECK",
        "",
        "本文件用于记录迁移重建后的系统状态快照，便于再次核对。",
        "",
        f"- 生效底稿链：{', '.join(row['report_date'] for row in active_drafts)}",
        "",
        "## 模块状态",
        "",
    ]
    for definition in SECTION_DEFINITIONS:
        if definition["key"] == "report_note":
            continue
        counts = by_module.get(definition["title"], {})
        summary = " / ".join(f"{label} {count}" for label, count in counts.items() if count) or "无结构化条目"
        lines.append(f"- {definition['number']}. {definition['title']}：{summary}")
    output_path.write_text("\n".join(lines), encoding="utf-8")
    return output_path


def build_linked_source_registry(_evidence_sources: list[dict]) -> list[dict]:
    registry: list[dict] = []
    for path in sorted(current_app.config["LINKED_DATA_ROOT"].glob("*.docx")):
        registry.extend(load_linked_docx_rows(path))
    return registry


def load_linked_docx_rows(path: Path) -> list[dict]:
    rows: list[dict] = []
    try:
        with open(path, "rb") as stream:
            document = Document(stream)
    except OSError:
        return rows

    for table in document.tables:
        normalized_rows = []
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                normalized_rows.append(values)
        if len(normalized_rows) < 2:
            continue
        headers = normalized_rows[0]
        mapped_headers = map_linked_headers(headers)
        if "title" not in mapped_headers.values() or "source_url" not in mapped_headers.values():
            continue
        for body_row in normalized_rows[1:]:
            record = {}
            for column_index, value in enumerate(body_row):
                key = mapped_headers.get(column_index)
                if key:
                    record[key] = value.strip()
            if not record.get("title"):
                continue
            record["source_level"] = normalize_source_level("", record.get("source_name", ""))
            record["event_key"] = provisional_event_key(
                "linked",
                record.get("title", ""),
                record.get("time_text", ""),
                record.get("core_content", ""),
                record.get("source_name", ""),
            )
            record["path"] = str(path)
            rows.append(record)
    return rows


def map_linked_headers(headers: list[str]) -> dict[int, str]:
    mapping: dict[int, str] = {}
    for index, header in enumerate(headers):
        normalized = normalize_compare_text(header)
        if "标题" in normalized:
            mapping[index] = "title"
        elif "时间" in normalized:
            mapping[index] = "time_text"
        elif "原始链接" in normalized or ("链接" in normalized and "原始" in normalized):
            mapping[index] = "source_url"
        elif "来源" in normalized:
            mapping[index] = "source_name"
        elif "核心内容" in normalized:
            mapping[index] = "core_content"
        elif "为什么" in normalized:
            mapping[index] = "why_included"
        elif "信息类型" in normalized:
            mapping[index] = "source_title"
    return mapping


def annotate_entry_timelines(entries: list[dict]) -> None:
    event_dates: defaultdict[str, list[str]] = defaultdict(list)
    for item in entries:
        if item.get("entry_type") == "placeholder" or item.get("is_deleted"):
            continue
        event_dates[item["event_key"]].append(item["report_date"])
    for item in entries:
        dates = sorted(set(event_dates.get(item["event_key"], [item["report_date"]])))
        item["first_seen_date"] = dates[0] if dates else item["report_date"]
        item["last_seen_date"] = dates[-1] if dates else item["report_date"]


def get_link_row_date(row: dict) -> str:
    event_day = extract_first_date(row.get("time_text", ""))
    return event_day.isoformat() if event_day else ""


def resolve_safe_link_candidates(item: dict, candidates: list[dict]) -> tuple[list[dict], str]:
    item_date = get_entry_match_date(item)
    item_source = get_entry_normalized_source(item)

    exact_steps = [
        (
            "title_date_source_exact",
            [
                row
                for row in candidates
                if item_date
                and item_source
                and get_link_row_date(row) == item_date
                and normalize_compare_text(row.get("source_name", "")) == item_source
            ],
        ),
        (
            "title_date_exact",
            [row for row in candidates if item_date and get_link_row_date(row) == item_date],
        ),
        (
            "title_source_exact",
            [
                row
                for row in candidates
                if item_source and normalize_compare_text(row.get("source_name", "")) == item_source
            ],
        ),
        ("title_exact_unique", candidates),
    ]
    for mode, rows in exact_steps:
        if len(rows) == 1:
            return rows, mode
        if len(rows) > 1:
            return [], f"ambiguous_{mode}"
    return [], "no_exact_match"


def resolve_legacy_fuzzy_link_candidates(item: dict, grouped_registry: dict[str, list[dict]]) -> tuple[list[dict], str]:
    if get_runtime_flag("LINKED_SOURCE_SAFE_MODE", True):
        return [], "safe_mode_exact_only"
    if not get_runtime_flag("ENABLE_LEGACY_FUZZY_LINK_BACKFILL", LEGACY_FUZZY_LINK_BACKFILL_ENABLED):
        return [], "legacy_fuzzy_disabled"

    current_title = get_entry_normalized_title(item)
    if not current_title:
        return [], "legacy_fuzzy_disabled"

    threshold = float(
        get_runtime_flag("LEGACY_FUZZY_LINK_BACKFILL_THRESHOLD", LEGACY_FUZZY_LINK_BACKFILL_THRESHOLD)
    )
    scored_candidates: list[tuple[float, dict]] = []
    for key, rows in grouped_registry.items():
        if not key:
            continue
        similarity = SequenceMatcher(None, current_title, key).ratio()
        if similarity >= threshold:
            for row in rows:
                scored_candidates.append((similarity, row))
    if not scored_candidates:
        return [], "legacy_fuzzy_none"

    scored_candidates.sort(
        key=lambda pair: (
            pair[0],
            get_link_row_date(pair[1]),
            ENTRY_SOURCE_LEVELS.get(pair[1].get("source_level", "C"), 1),
        ),
        reverse=True,
    )
    best_score = scored_candidates[0][0]
    best_rows = [row for score, row in scored_candidates if score == best_score]
    if len(best_rows) == 1:
        return best_rows, "legacy_fuzzy_title"
    return [], "ambiguous_legacy_fuzzy_title"


def find_linked_match_result(item: dict, grouped_registry: dict[str, list[dict]]) -> tuple[list[dict], str]:
    current_title = get_entry_normalized_title(item)
    if not current_title:
        return [], "missing_title"
    exact_candidates = list(grouped_registry.get(current_title, []))
    matches, mode = resolve_safe_link_candidates(item, exact_candidates)
    if matches or mode.startswith("ambiguous_"):
        return matches, mode
    return resolve_legacy_fuzzy_link_candidates(item, grouped_registry)


def enrich_entries_with_linked_sources(entries: list[dict], linked_registry: list[dict]) -> None:
    if not linked_registry:
        return
    grouped_registry: defaultdict[str, list[dict]] = defaultdict(list)
    for record in linked_registry:
        grouped_registry[normalize_compare_text(record.get("title", ""))].append(record)

    for item in entries:
        if item.get("entry_type") == "placeholder":
            continue
        item.setdefault("evidence_json", {})
        if item.get("source_url"):
            item["evidence_json"]["link_backfill_mode"] = "preserved_existing_source_url"
            item["evidence_json"]["match_mode"] = "existing_source_url"
            continue
        matches, match_mode = find_linked_match_result(item, grouped_registry)
        if not matches:
            item["evidence_json"]["link_backfill_mode"] = (
                "skipped_ambiguous" if match_mode.startswith("ambiguous_") else "no_match"
            )
            item["evidence_json"]["match_mode"] = match_mode
            continue
        primary = matches[0]
        item["source_url"] = primary.get("source_url", "")
        if not item.get("source_title"):
            item["source_title"] = primary.get("title", "")
        item["evidence_json"]["linked_match_count"] = len(matches)
        item["evidence_json"]["linked_primary_path"] = primary.get("path", "")
        item["evidence_json"]["match_mode"] = match_mode
        item["evidence_json"]["link_backfill_mode"] = (
            "legacy_fuzzy_backfill" if match_mode.startswith("legacy_fuzzy_") else "exact_safe_first"
        )
        supporting = []
        for match in matches[:5]:
            supporting.append(
                {
                    "title": match.get("title", ""),
                    "source_name": match.get("source_name", ""),
                    "source_url": match.get("source_url", ""),
                    "time_text": match.get("time_text", ""),
                }
            )
        item["supporting_sources_json"] = supporting


def find_linked_matches(item: dict, grouped_registry: dict[str, list[dict]]) -> list[dict]:
    matches, _match_mode = find_linked_match_result(item, grouped_registry)
    return matches


def join_evidence_names(items: list[dict]) -> str:
    names = [item["name"] for item in items if item]
    return "；".join(names) if names else "当前目录证据集"


def section_type_from_path(path: str) -> str:
    if "." not in path:
        return ""
    return SECTION_TYPE_BY_SUFFIX.get(path.split(".")[-1], "")


def canonical_subsection_path(path: str) -> str:
    parts = [piece for piece in path.split(".") if piece]
    if len(parts) >= 2:
        return f"{parts[0]}.{parts[1]}"
    return path


def expected_subsection_title(path: str) -> str:
    suffix = path.split(".")[-1] if "." in path else ""
    for code, title in DRAFT_STANDARD_SUBSECTIONS:
        if code == suffix:
            return title
    return path


def trim_row(row: list[str]) -> list[str]:
    last_non_empty = 0
    for index, value in enumerate(row, start=1):
        if (value or "").strip():
            last_non_empty = index
    clipped = row[:last_non_empty] if last_non_empty else row
    return [cell.strip() for cell in clipped]


def normalize_business_tags(raw_value) -> list[str]:
    if isinstance(raw_value, list):
        values = [str(item).strip() for item in raw_value if str(item).strip()]
    else:
        values = [part.strip() for part in re.split(r"[、,，/｜|；;\n]+", str(raw_value or "")) if part.strip()]
    seen: set[str] = set()
    normalized: list[str] = []
    for value in values:
        key = BUSINESS_TAG_CANONICAL_MAP.get(normalize_compare_text(value)) or normalize_compare_text(value)
        if not key or key in seen:
            continue
        seen.add(key)
        normalized.append(key)
    return normalized[:8]


def looks_like_placeholder_text(text: str) -> bool:
    normalized = normalize_compare_text(text)
    return any(normalize_compare_text(pattern) in normalized for pattern in PLACEHOLDER_PATTERNS)


def looks_like_placeholder_row(mapped: dict) -> bool:
    if looks_like_placeholder_text(mapped.get("title", "")):
        return True
    if looks_like_placeholder_text(mapped.get("core_content", "")):
        return True
    meaningful = []
    for key, value in mapped.items():
        if key in {"title", "core_content"} or not value:
            continue
        if isinstance(value, list):
            if any(str(item).strip() for item in value):
                meaningful.append(value)
            continue
        stripped = str(value).strip()
        if stripped and stripped.strip("/") and stripped.strip("—"):
            meaningful.append(value)
    return not meaningful and looks_like_placeholder_text(f"{mapped.get('title', '')} {mapped.get('core_content', '')}")


def normalize_source_level(level_text: str, source_name: str) -> str:
    normalized = (level_text or "").upper().replace("级", "").strip()
    if normalized in ENTRY_SOURCE_LEVELS:
        return normalized
    source = source_name or ""
    if any(pattern in source for pattern in OFFICIAL_SOURCE_PATTERNS):
        return "A2"
    if has_weak_source(source):
        return "C"
    if source:
        return "B"
    return "C"


def infer_confidence_level(source_level: str, needs_review: bool) -> str:
    if source_level in {"A1", "A2"} and not needs_review:
        return "高"
    if source_level == "占位":
        return "低"
    if source_level == "C" or needs_review:
        return "低"
    return "中"


def provisional_event_key(
    module_key: str,
    title: str,
    time_text: str,
    core_content: str,
    source_name: str,
    *,
    canonical_title: str = "",
) -> str:
    event_day = extract_first_date(time_text)
    event_date = event_day.isoformat() if event_day else ""
    stable_title = normalize_compare_text(canonical_title) or normalize_compare_text(title) or normalize_compare_text(core_content)
    stable_source = normalize_compare_text(source_name)
    fingerprint = "|".join([module_key, stable_title, event_date, stable_source])
    digest = hashlib.sha1(fingerprint.encode("utf-8")).hexdigest()[:16]
    return f"{module_key}:{digest}"


def is_date_in_window(event_day: date | None, window: tuple[date, date]) -> bool:
    if not event_day:
        return False
    return window[0] <= event_day <= window[1]


def extract_source_url(source_name: str) -> str:
    match = SOURCE_URL_PATTERN.search(source_name or "")
    return match.group(0) if match else ""


def has_weak_source(source_name: str) -> bool:
    source = source_name or ""
    return any(pattern in source for pattern in WEAK_SOURCE_PATTERNS)


def is_missing_minimum_fields(item: dict) -> bool:
    if not (item.get("title") or "").strip():
        return True
    if not (item.get("core_content") or "").strip():
        return True
    if item["section_type"] in {"新增", "重点"} and not (item.get("time_text") or "").strip():
        return True
    if item["section_type"] in {"新增", "重点"} and not (item.get("source_name") or "").strip():
        return True
    return False


def find_previous_match(previous_items: list[dict], current_item: dict) -> dict | None:
    exact_match = find_exact_event_match(previous_items, current_item)
    if exact_match:
        return exact_match

    fuzzy_enabled = get_runtime_flag("ENABLE_LEGACY_V1_FUZZY_MATCH", LEGACY_V1_FUZZY_MATCH_ENABLED)
    if not fuzzy_enabled or has_delta_capability(current_item):
        return None

    threshold = float(get_runtime_flag("LEGACY_V1_FUZZY_MATCH_THRESHOLD", LEGACY_V1_FUZZY_MATCH_THRESHOLD))
    candidates = []
    for item in iter_real_candidates(previous_items):
        score = event_similarity(item, current_item)
        if score >= threshold:
            candidates.append((score, item))
    if not candidates:
        return None
    candidates.sort(key=lambda pair: (pair[0], pair[1].get("report_date", "")), reverse=True)
    return candidates[0][1]


def event_similarity(left: dict, right: dict) -> float:
    title_score = SequenceMatcher(None, normalize_compare_text(left["title"]), normalize_compare_text(right["title"])).ratio()
    core_score = SequenceMatcher(None, normalize_compare_text(left["core_content"]), normalize_compare_text(right["core_content"])).ratio()
    same_day_bonus = 0.08 if left.get("event_date") and left.get("event_date") == right.get("event_date") else 0.0
    same_source_bonus = 0.05 if normalize_compare_text(left.get("source_name", "")) == normalize_compare_text(right.get("source_name", "")) else 0.0
    return max(title_score * 0.7 + core_score * 0.3 + same_day_bonus + same_source_bonus, core_score)


def get_entry_evidence(item: dict) -> dict:
    evidence = item.get("evidence_json") or {}
    if isinstance(evidence, str):
        return load_json(evidence, {})
    return evidence


def get_entry_contract_version(item: dict) -> str:
    return str(get_entry_evidence(item).get("contract_version") or "").strip().lower()


def get_entry_delta_text(item: dict) -> str:
    return str(get_entry_evidence(item).get("delta_text") or "").strip()


def get_entry_business_tags(item: dict) -> list[str]:
    return normalize_business_tags(get_entry_evidence(item).get("business_tags"))


def get_runtime_flag(name: str, default):
    if has_app_context():
        return current_app.config.get(name, default)
    return default


def get_entry_event_anchor(item: dict) -> str:
    return normalize_compare_text(get_entry_evidence(item).get("event_anchor", ""))


def get_entry_canonical_title(item: dict) -> str:
    return normalize_compare_text(get_entry_evidence(item).get("canonical_title", ""))


def get_entry_normalized_title(item: dict) -> str:
    return normalize_compare_text(item.get("title", ""))


def get_entry_normalized_source(item: dict) -> str:
    return normalize_compare_text(item.get("source_name", ""))


def get_entry_match_date(item: dict) -> str:
    event_date = (item.get("event_date") or "").strip()
    if event_date:
        return event_date
    event_day = extract_first_date(item.get("time_text", ""))
    return event_day.isoformat() if event_day else ""


def pick_best_event_candidate(candidates: list[dict]) -> dict | None:
    if not candidates:
        return None
    return max(
        candidates,
        key=lambda item: (
            item.get("report_date", ""),
            item.get("last_seen_date", ""),
            ENTRY_SOURCE_LEVELS.get(item.get("source_level", "C"), 1),
            len(normalize_compare_text(item.get("core_content", ""))),
        ),
    )


def iter_real_candidates(items: list[dict]) -> list[dict]:
    return [item for item in items if item.get("entry_type") == "real" and not item.get("is_deleted")]


def find_exact_event_match(previous_items: list[dict], current_item: dict) -> dict | None:
    candidates = iter_real_candidates(previous_items)
    if not candidates:
        return None

    current_anchor = get_entry_event_anchor(current_item)
    if current_anchor:
        anchor_matches = [item for item in candidates if get_entry_event_anchor(item) == current_anchor]
        if anchor_matches:
            return pick_best_event_candidate(anchor_matches)

    current_canonical = get_entry_canonical_title(current_item)
    if current_canonical:
        canonical_matches = [item for item in candidates if get_entry_canonical_title(item) == current_canonical]
        if canonical_matches:
            return pick_best_event_candidate(canonical_matches)

    current_title = get_entry_normalized_title(current_item)
    current_date = get_entry_match_date(current_item)
    current_source = get_entry_normalized_source(current_item)
    if not current_title:
        return None

    title_matches = [item for item in candidates if get_entry_normalized_title(item) == current_title]
    if current_date and current_source:
        title_date_source_matches = [
            item
            for item in title_matches
            if get_entry_match_date(item) == current_date and get_entry_normalized_source(item) == current_source
        ]
        if title_date_source_matches:
            return pick_best_event_candidate(title_date_source_matches)

    if current_date:
        title_date_matches = [item for item in title_matches if get_entry_match_date(item) == current_date]
        if title_date_matches:
            return pick_best_event_candidate(title_date_matches)

    if current_source:
        title_source_matches = [
            item for item in title_matches if get_entry_normalized_source(item) == current_source
        ]
        if title_source_matches:
            return pick_best_event_candidate(title_source_matches)

    if len(title_matches) == 1:
        return title_matches[0]
    return None


def has_delta_capability(item: dict) -> bool:
    evidence = get_entry_evidence(item)
    contract_version = get_entry_contract_version(item)
    if contract_version:
        return contract_version == "v2"
    return bool(str(evidence.get("delta_text") or "").strip())


def has_new_time_node(previous_item: dict, current_item: dict) -> bool:
    previous_dates = {value.isoformat() for value in extract_all_dates(previous_item.get("time_text", ""))}
    current_dates = {value.isoformat() for value in extract_all_dates(current_item.get("time_text", ""))}
    return bool(current_dates - previous_dates)


def has_source_url_gain(previous_item: dict, current_item: dict) -> bool:
    return not (previous_item.get("source_url") or "").strip() and bool((current_item.get("source_url") or "").strip())


def has_meaningful_delta(previous_item: dict, current_item: dict) -> bool:
    delta_text = get_entry_delta_text(current_item)
    source_gain = ENTRY_SOURCE_LEVELS.get(current_item.get("source_level", "C"), 1) > ENTRY_SOURCE_LEVELS.get(previous_item.get("source_level", "C"), 1)
    source_url_gain = has_source_url_gain(previous_item, current_item)
    new_time_node = has_new_time_node(previous_item, current_item)
    if delta_text:
        return True
    if source_gain or source_url_gain or new_time_node:
        return True
    if has_delta_capability(previous_item) or has_delta_capability(current_item):
        return False

    core_similarity = SequenceMatcher(
        None,
        normalize_compare_text(previous_item.get("core_content", "")),
        normalize_compare_text(current_item.get("core_content", "")),
    ).ratio()
    why_similarity = SequenceMatcher(
        None,
        normalize_compare_text(previous_item.get("why_included", "")),
        normalize_compare_text(current_item.get("why_included", "")),
    ).ratio()
    if core_similarity < 0.82:
        return True
    return core_similarity < 0.9 and why_similarity < 0.45


def has_overclaimed_value(item: dict) -> bool:
    why_text = item.get("why_included", "")
    core_text = item.get("core_content", "")
    if not why_text:
        return False
    if any(token in why_text for token in ["行业趋势", "全国性", "政策导向"]) and all(token not in core_text for token in ["政策", "官方", "政府", "协会", "发布"]):
        return True
    return False


def entry_quality_score(item: dict) -> tuple:
    return (
        1 if not item["is_deleted"] else 0,
        ENTRY_SOURCE_LEVELS.get(item.get("source_level", "C"), 1),
        0 if item.get("needs_review") else 1,
        len(normalize_compare_text(item.get("core_content", ""))),
        item.get("report_date", ""),
    )


def find_canonical_group(groups: list[dict], current_item: dict) -> dict | None:
    return find_previous_match(groups, current_item)


def same_effective_fact(left: dict, right: dict) -> bool:
    if left.get("event_key") and right.get("event_key") and left.get("event_key") != right.get("event_key"):
        return False
    if has_meaningful_delta(left, right) or has_meaningful_delta(right, left):
        return False
    left_delta = normalize_compare_text(get_entry_delta_text(left))
    right_delta = normalize_compare_text(get_entry_delta_text(right))
    if left_delta != right_delta and (left_delta or right_delta):
        return False
    core_similarity = SequenceMatcher(
        None,
        normalize_compare_text(left.get("core_content", "")),
        normalize_compare_text(right.get("core_content", "")),
    ).ratio()
    merged_length_gap = abs(len(normalize_compare_text(left["core_content"])) - len(normalize_compare_text(right["core_content"])))
    return core_similarity >= 0.9 and merged_length_gap <= 24


def merge_duplicate_into_canonical(canonical: dict, duplicate: dict) -> None:
    canonical_source_level = ENTRY_SOURCE_LEVELS.get(canonical.get("source_level", "C"), 1)
    duplicate_source_level = ENTRY_SOURCE_LEVELS.get(duplicate.get("source_level", "C"), 1)
    if duplicate_source_level > canonical_source_level:
        canonical["source_level"] = duplicate["source_level"]
        canonical["source_name"] = duplicate["source_name"]
        canonical["source_url"] = duplicate["source_url"]
        canonical["confidence_level"] = infer_confidence_level(canonical["source_level"], bool(canonical.get("needs_review")))
    canonical_evidence = deepcopy(get_entry_evidence(canonical))
    duplicate_evidence = get_entry_evidence(duplicate)
    canonical_evidence["delta_text"] = merge_text(
        canonical_evidence.get("delta_text", ""),
        duplicate_evidence.get("delta_text", ""),
    )
    merged_tags = normalize_business_tags(get_entry_business_tags(canonical) + get_entry_business_tags(duplicate))
    canonical_evidence["business_tags"] = merged_tags
    if not canonical.get("source_url") and duplicate.get("source_url"):
        canonical["source_url"] = duplicate["source_url"]
    canonical["core_content"] = merge_text(canonical["core_content"], duplicate["core_content"])
    canonical["why_included"] = merge_text(canonical.get("why_included", ""), duplicate.get("why_included", ""))
    if duplicate.get("needs_review"):
        canonical["needs_review"] = 1
    canonical["evidence_json"] = canonical_evidence
    canonical["updated_at"] = now_string()


def merge_text(existing: str, incoming: str) -> str:
    existing = (existing or "").strip()
    incoming = (incoming or "").strip()
    if not incoming:
        return existing
    if not existing:
        return incoming
    existing_norm = normalize_compare_text(existing)
    incoming_norm = normalize_compare_text(incoming)
    if incoming_norm in existing_norm:
        return existing
    if existing_norm in incoming_norm:
        return incoming
    return f"{existing}\n补充：{incoming}"


def filter_current_group_rows(rows: list[dict], section_type: str) -> list[dict]:
    target_rows = [row for row in rows if row["section_type"] == section_type and not row["is_deleted"]]
    real_rows = [row for row in target_rows if row["entry_type"] == "real" and row["display_status"] in {"新增", "更新"}]
    if real_rows:
        return real_rows
    placeholder_rows = [row for row in target_rows if row["entry_type"] == "placeholder"]
    return placeholder_rows[:1]


def latest_effective_rows(rows: list[dict], section_type: str) -> list[dict]:
    relevant = [row for row in rows if row["section_type"] == section_type and not row["is_deleted"]]
    canonical: dict[str, dict] = {}
    for row in relevant:
        if row["display_status"] not in {"背景补充", "历史保留", "新增", "更新"}:
            continue
        if row["event_key"] not in canonical or entry_quality_score(row) > entry_quality_score(canonical[row["event_key"]]):
            canonical[row["event_key"]] = row
    values = list(canonical.values())
    values.sort(key=lambda item: (item["report_date"], item["title"]))
    return values


def latest_note_rows(rows: list[dict]) -> list[dict]:
    note_rows = [row for row in rows if row["section_type"] == "当前状态" and not row["is_deleted"]]
    if not note_rows:
        return []
    latest_date = max(row["report_date"] for row in note_rows)
    return [row for row in note_rows if row["report_date"] == latest_date]


def canonical_historical_rows(rows: list[dict], excluded_event_keys: set[str]) -> list[dict]:
    canonical: dict[str, dict] = {}
    for row in rows:
        if row["entry_type"] != "real" or row["is_deleted"]:
            continue
        if row["section_type"] == "当前状态":
            continue
        if row["event_key"] in excluded_event_keys:
            continue
        if row["display_status"] not in {"新增", "更新", "历史保留", "背景补充"}:
            continue
        chosen = canonical.get(row["event_key"])
        if not chosen or entry_quality_score(row) > entry_quality_score(chosen):
            retained = deepcopy(row)
            retained["display_status"] = "历史保留"
            canonical[row["event_key"]] = retained
    values = list(canonical.values())
    values.sort(key=lambda item: (item["report_date"], item["title"]))
    return values


def resolve_display_category_from_text(section_key: str, text: str) -> str:
    normalized = normalize_compare_text(text or "")
    if not normalized:
        return ""
    for rule in SECTION_SUBCATEGORY_RULES.get(section_key, []):
        rule_title = rule.get("title", "")
        if normalize_compare_text(rule_title) == normalized:
            return rule_title
        for keyword in rule.get("keywords", []):
            normalized_keyword = normalize_compare_text(keyword)
            if normalized_keyword and normalized_keyword in normalized:
                return rule_title
    return ""


def resolve_display_category_from_tags(section_key: str, tags: list[str]) -> str:
    normalized_tags = {
        BUSINESS_TAG_CANONICAL_MAP.get(normalize_compare_text(tag)) or normalize_compare_text(tag)
        for tag in tags
        if normalize_compare_text(tag)
    }
    if not normalized_tags:
        return ""
    for rule in SECTION_SUBCATEGORY_RULES.get(section_key, []):
        rule_tokens = {
            BUSINESS_TAG_CANONICAL_MAP.get(normalize_compare_text(value)) or normalize_compare_text(value)
            for value in [rule.get("key", ""), *rule.get("tag_keys", [])]
            if normalize_compare_text(value)
        }
        if rule_tokens & normalized_tags:
            return rule.get("title", "")
    return ""


def resolve_row_display_category(section_key: str, row: dict) -> str:
    resolved = resolve_display_category_from_tags(section_key, get_entry_business_tags(row))
    if resolved:
        return resolved
    fallback_text = " ".join(
        part
        for part in [
            row.get("title", ""),
            row.get("core_content", ""),
            row.get("why_included", ""),
            row.get("source_name", ""),
        ]
        if part
    )
    return resolve_display_category_from_text(section_key, fallback_text)


def build_display_entry_groups(section_key: str, title: str, level: int, rows: list[dict], category: str) -> list[dict]:
    if not rows:
        return []
    if section_key not in SECTION_SUBCATEGORY_RULES:
        return [build_entry_group(title, level, rows, category)]

    grouped_rows: defaultdict[str, list[dict]] = defaultdict(list)
    ungrouped_rows: list[dict] = []
    for row in rows:
        if row.get("entry_type") != "real":
            ungrouped_rows.append(row)
            continue
        display_category = resolve_row_display_category(section_key, row)
        if display_category:
            grouped_rows[display_category].append(row)
        else:
            ungrouped_rows.append(row)

    groups: list[dict] = []
    for rule in SECTION_SUBCATEGORY_RULES.get(section_key, []):
        group_title = rule.get("title", "")
        if grouped_rows.get(group_title):
            groups.append(build_entry_group(title, level, grouped_rows[group_title], group_title))
    if ungrouped_rows:
        groups.append(build_entry_group(title, level, ungrouped_rows, category))
    return groups


def build_entry_group(title: str, level: int, rows: list[dict], category: str) -> dict:
    blocks: list[dict] = []
    for row in rows:
        card = entry_to_card(row)
        blocks.append({"type": "card", "card": card})
    return {"title": title, "level": level, "category": category, "blocks": blocks}


def build_note_group(title: str, rows: list[dict]) -> dict:
    blocks = [{"type": "paragraph", "text": row["core_content"]} for row in rows if row.get("core_content")]
    return {"title": title, "level": 2, "category": "当前状态", "blocks": blocks}


def entry_to_card(row: dict) -> dict:
    business_tags = get_entry_business_tags(row)
    delta_text = get_entry_delta_text(row)
    tags = [row["source_level"]]
    if row["section_type"] in {"新增", "重点", "背景"}:
        tags.append(row["section_type"])
    if row.get("needs_review"):
        tags.append("低可信待复核")
    if row["entry_type"] == "placeholder":
        tags = ["占位项"]
    return build_card(
        title=row["title"],
        status="待人工复核" if row.get("needs_review") and row["display_status"] != "占位项" else row["display_status"],
        time=row.get("time_text") or None,
        source=row.get("source_name") or None,
        core_content=row.get("core_content") or row.get("note_text") or row["title"],
        why=row.get("why_included") or None,
        tags=tags,
        style="intelligence",
        source_url=row.get("source_url") or None,
        delta_text=delta_text or None,
        business_tags=business_tags,
        source_title=row.get("source_title") or row.get("title") or None,
        supporting_sources=row.get("supporting_sources_json") or [],
        first_seen=row.get("first_seen_date") or row.get("report_date") or None,
        last_seen=row.get("last_seen_date") or row.get("report_date") or None,
        confidence_level=row.get("confidence_level") or None,
        needs_review=bool(row.get("needs_review")),
        entry_id=row["id"],
        compare_meta={
            "title": row["title"],
            "time": row.get("time_text", ""),
            "source": row.get("source_name", ""),
            "source_url": row.get("source_url", ""),
            "core_content": row.get("core_content", ""),
            "why": row.get("why_included", ""),
            "object_name": row["title"],
            "event_key": row["event_key"],
            "delta_text": delta_text,
            "business_tags": business_tags,
        },
    )


def count_group_statuses(groups: list[dict]) -> dict:
    counts = {label: 0 for label in DISPLAY_STATUS_ORDER}
    for group in groups:
        for block in group.get("blocks", []):
            if block.get("type") != "card":
                continue
            card = block["card"]
            status = card.get("status")
            if status == "待人工复核":
                status = "更新"
            if status not in counts:
                counts[status] = 0
            counts[status] += 1
    return counts


def section_status_from_groups(groups: list[dict], counts: dict) -> str:
    if counts.get("新增"):
        return "新增"
    if counts.get("更新"):
        return "更新"
    if counts.get("背景补充"):
        return "背景补充"
    if counts.get("历史保留"):
        return "历史保留"
    if counts.get("占位项"):
        return "占位项"
    return "无内容"


def section_status_from_counts(counts: dict) -> str:
    if counts.get("新增"):
        return "新增"
    if counts.get("更新"):
        return "更新"
    if counts.get("背景补充"):
        return "背景补充"
    if counts.get("历史保留"):
        return "历史保留"
    if counts.get("占位项"):
        return "占位项"
    return "无内容"


def build_rollup_note_from_entries(report_date: str, counts: dict, source_dates: list[str]) -> str:
    parts = [f"{label} {count} 条" for label, count in counts.items() if count and label != "无内容"]
    if not source_dates:
        return "本板块当前没有可累计显示的有效内容。"
    source_range = source_dates[0] if len(source_dates) == 1 else f"{source_dates[0]} 至 {source_dates[-1]}"
    if parts:
        return f"截至 {report_date} 的累计有效内容已重建完成，覆盖来源日期 {source_range}；当前展示为 {' / '.join(parts)}。"
    return f"截至 {report_date} 的累计有效内容链已重建，但本板块仅保留说明或空结果结构。"


def build_history_note(counts: dict) -> str:
    parts = [f"{label} {count} 条" for label, count in counts.items() if count]
    return " / ".join(parts) if parts else "本日无有效结构化内容"


def render_text_from_groups(groups: list[dict]) -> str:
    parts: list[str] = []
    for group in groups:
        if group.get("title"):
            parts.append(group["title"])
        for block in group.get("blocks", []):
            if block.get("type") == "paragraph":
                parts.append(block.get("text", ""))
            elif block.get("type") == "card":
                card = block["card"]
                parts.extend([card.get("title", ""), card.get("core_content", ""), card.get("why", "")])
    return "\n\n".join(part for part in parts if part).strip()


def build_decision(item: dict, action: str, reason: str) -> dict:
    return {
        "report_date": item["report_date"],
        "module_name": item["module_name"],
        "title": item["title"],
        "action": action,
        "reason": reason,
    }
